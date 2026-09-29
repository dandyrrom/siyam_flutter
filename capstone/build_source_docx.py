"""Builds the printable SIYAM source-code listing (.docx).

Usage (from the repo root):
    python capstone/build_source_docx.py [--target-pages 800] [--out PATH]

Page setup: Letter, margins 1" top/right/bottom and 1.5" left, Times New Roman
10 pt, exact 11.5 pt line spacing (so lines per page are predictable).
Runs of blank lines are collapsed to one; comments are kept.

Pages are counted by rendering a throwaway PDF with LibreOffice (soffice) and
poppler (pdfinfo / pdftotext). The PDF is never written next to the output.
"""

import argparse
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
FONT = "Times New Roman"
FONT_SIZE = Pt(10)
LINE_HEIGHT = Pt(11.5)
INDENT_PER_SPACE_TWIPS = 60
WRAP_HANGING_TWIPS = 240
MAX_INDENT_TWIPS = 4400
TEXT_WIDTH = Inches(6)

# Screens are the only files that get trimmed. They are added in this order
# (most important first) until the page budget is used up. Whole files only.
PAGE_PRIORITY = [
    "lib/pages/login_page.dart",
    "lib/pages/register_page.dart",
    "lib/pages/dashboard_page.dart",
    "lib/pages/dashboard/manager_dashboard.dart",
    "lib/pages/dashboard/staff_dashboard.dart",
    "lib/pages/dashboard/donor_dashboard.dart",
    "lib/pages/inventory_page.dart",
    "lib/pages/inventory_item_page.dart",
    "lib/pages/add_item_page.dart",
    "lib/pages/add_treatment_page.dart",
    "lib/pages/treatment_detail_page.dart",
    "lib/pages/medical_records_page.dart",
    "lib/pages/donations_page.dart",
    "lib/pages/donor/donate_page.dart",
    "lib/pages/donor/donation_history_page.dart",
    "lib/pages/donor/impacts_page.dart",
    "lib/pages/submission_detail_page.dart",
    "lib/pages/purchase_orders_page.dart",
    "lib/pages/purchase_trans_page.dart",
    "lib/pages/suppliers_page.dart",
    "lib/pages/audit_trail_page.dart",
    "lib/pages/animal_records_page.dart",
    "lib/pages/animal_medical_history_page.dart",
    "lib/pages/notifications_page.dart",
    "lib/pages/notification_detail_page.dart",
    "lib/pages/profile_page.dart",
    "lib/pages/reports_page.dart",
    "lib/pages/reports/manager_reports_page.dart",
    "lib/pages/reports/staff_reports_page.dart",
    "lib/pages/settings_page.dart",
]

# Not in use (nothing imports them) - never printed.
UNUSED_FILES = {
    "lib/pages/placeholder_page.dart",
    "lib/pages/replenishment_page.dart",
}


def dart_files(pattern):
    return sorted(
        p.relative_to(ROOT).as_posix()
        for p in ROOT.glob(pattern)
        if p.is_file() and p.relative_to(ROOT).as_posix() not in UNUSED_FILES
    )


def fixed_groups():
    return [
        ("Application Entry Point", ["lib/main.dart", "lib/app.dart"]),
        ("Core (theme, colors, validators, config)", dart_files("lib/core/*.dart")),
        ("Routing and Navigation", dart_files("lib/routing/*.dart")),
        ("Data Models", dart_files("lib/models/*.dart")),
        ("State Management", dart_files("lib/state/*.dart")),
        ("Services (business logic)", dart_files("lib/services/*.dart")),
        ("Mock Data Layer", dart_files("lib/mock/*.dart")),
        ("Shared Widgets", dart_files("lib/widgets/*.dart")),
    ]


def test_files():
    return dart_files("test/**/*.dart")


def read_source(rel):
    text = (ROOT / rel).read_text(encoding="utf-8", errors="replace")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text.replace("\r\n", "\n").replace("\r", "\n"))
    text = text.replace("\t", "    ")
    out, prev_blank = [], True
    for line in text.split("\n"):
        line = line.rstrip()
        blank = not line
        if blank and prev_blank:
            continue
        out.append(line)
        prev_blank = blank
    while out and not out[-1]:
        out.pop()
    return out


def build_groups(selected_pages):
    groups = fixed_groups()
    pages_sel = [f for f in PAGE_PRIORITY if f in selected_pages]
    page_order = sorted(pages_sel, key=lambda f: (f.count("/"), f))
    groups.append(("Screens (pages)", page_order))
    groups.append(("Unit Tests", test_files()))
    return groups


def omitted_pages(selected_pages):
    all_pages = set(dart_files("lib/pages/**/*.dart"))
    return sorted(all_pages - set(selected_pages))


def set_base_style(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = FONT_SIZE
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = LINE_HEIGHT
    pf.widow_control = False
    ppr = normal.element.get_or_add_pPr()
    spacing = ppr.find(qn("w:spacing"))
    spacing.set(qn("w:lineRule"), "exact")


def set_page_setup(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.right_margin = Inches(1)
    section.left_margin = Inches(1.5)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def setup_header_footer(section):
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.text = "Source Code Listing"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE")


def add_bottom_border(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    borders.append(bottom)
    ppr.append(borders)


def code_paragraph(body, line):
    p = OxmlElement("w:p")
    stripped = line.lstrip(" ")
    lead = len(line) - len(stripped)
    hanging = WRAP_HANGING_TWIPS
    left = min(INDENT_PER_SPACE_TWIPS * lead, MAX_INDENT_TWIPS) + hanging
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    ppr.append(ind)
    p.append(ppr)
    if stripped:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = stripped
        r.append(t)
        p.append(r)
    body.insert(len(body) - 1, p)


def bold_run(paragraph, text, size=None):
    run = paragraph.add_run(text)
    run.bold = True
    if size:
        run.font.size = size
    return run


def add_title_page(doc):
    def line(text, size=10, bold=False, before=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        return p

    line("[INSTITUTION / SCHOOL NAME]", 14, True, 60)
    line("[College / Department / Program]", 12, False, 6)
    line("[PROJECT TITLE]", 24, True, 150)
    line("[Project subtitle or system description]", 14, False, 12)
    line("Source Code Listing", 18, True, 70)
    line("Submitted as a requirement for [Capstone Course Code and Title]", 12, False, 70)
    line("[Team Member Names]", 12, False, 40)
    line("[Adviser Name]", 12, False, 8)
    line("[Month Year]", 12, False, 40)


def add_toc(doc, groups, appendix_title, pages):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bold_run(p, "TABLE OF CONTENTS", Pt(14))

    def entry(text, page, bold=False, indent=0, before=0):
        para = doc.add_paragraph()
        pf = para.paragraph_format
        pf.left_indent = Inches(indent)
        pf.space_before = Pt(before)
        pf.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = para.add_run(f"{text}\t{page}")
        run.bold = bold

    for n, (title, files) in enumerate(groups, 1):
        if not files:
            continue
        entry(f"Part {n}. {title}", pages.get(("group", n), 0), bold=True, before=6)
        for f in files:
            entry(f, pages.get(f, 0), indent=0.3)
    entry(appendix_title, pages.get(("group", "appendix"), 0), bold=True, before=6)


def add_group_heading(doc, n, title):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(6)
    bold_run(p, f"PART {n}. {title.upper()}", Pt(12))
    add_bottom_border(p)
    return p


def add_file_heading(doc, rel):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    bold_run(p, f"Listing: {rel}")
    add_bottom_border(p)


def build_document(selected_pages, pages, out_path):
    groups = build_groups(selected_pages)
    appendix_title = "Appendix. Screen Files Not Printed in This Listing"
    doc = Document()
    set_base_style(doc)
    section = doc.sections[0]
    set_page_setup(section)
    setup_header_footer(section)

    add_title_page(doc)
    add_toc(doc, groups, appendix_title, pages)

    body = doc.element.body
    line_count = 0
    for n, (title, files) in enumerate(groups, 1):
        if not files:
            continue
        add_group_heading(doc, n, title)
        for rel in files:
            add_file_heading(doc, rel)
            lines = read_source(rel)
            line_count += len(lines)
            for line in lines:
                code_paragraph(body, line)

    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(6)
    bold_run(p, appendix_title.upper(), Pt(12))
    add_bottom_border(p)
    note = doc.add_paragraph(
        "The files below are part of the application but were left out of this printed "
        "listing to stay within the page limit. Each is a screen (page) file under lib/pages/."
    )
    note.paragraph_format.space_after = Pt(6)
    note.paragraph_format.line_spacing = 1.0
    for f in omitted_pages(selected_pages):
        doc.add_paragraph(f)

    doc.save(out_path)
    return groups, line_count


def render_pdf(docx_path, workdir):
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(workdir), str(docx_path)],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return Path(workdir) / (Path(docx_path).stem + ".pdf")


def pdf_pages(pdf):
    out = subprocess.run(["pdfinfo", str(pdf)], check=True, capture_output=True, text=True).stdout
    return int(re.search(r"Pages:\s+(\d+)", out).group(1))


def locate_headings(pdf, groups):
    text = subprocess.run(
        ["pdftotext", "-layout", str(pdf), "-"], check=True, capture_output=True, text=True
    ).stdout
    pages = {}
    for page_no, page_text in enumerate(text.split("\f"), 1):
        for line in page_text.splitlines():
            line = line.strip()
            m = re.match(r"Listing: (\S+)$", line)
            if m and m.group(1) not in pages:
                pages[m.group(1)] = page_no
            m = re.match(r"PART (\d+)\. ", line)
            if m and ("group", int(m.group(1))) not in pages:
                pages[("group", int(m.group(1)))] = page_no
            if line.startswith("APPENDIX. SCREEN FILES") and ("group", "appendix") not in pages:
                pages[("group", "appendix")] = page_no
    return pages


def count_lines(files):
    return sum(len(read_source(f)) + 4 for f in files)


def choose_pages(capacity_lines):
    chosen, used = [], 0
    for f in PAGE_PRIORITY:
        cost = count_lines([f])
        if used + cost <= capacity_lines:
            chosen.append(f)
            used += cost
    return chosen


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--target-pages", type=int, default=800)
    ap.add_argument("--out", default=str(ROOT / "capstone" / "SIYAM_Source_Code_800_pages.docx"))
    ap.add_argument("--capacity-lines", type=int, default=None)
    args = ap.parse_args()

    fixed = [f for _, files in fixed_groups() for f in files] + test_files()
    fixed_lines = count_lines(fixed)
    lines_per_page = 55.0
    capacity = args.capacity_lines or int(args.target_pages * lines_per_page - fixed_lines - 400)

    tmp = Path(tempfile.mkdtemp(prefix="siyam_docx_"))
    best = None
    for attempt in range(1, 7):
        selected = choose_pages(capacity)
        draft = tmp / "draft.docx"
        groups, total_lines = build_document(selected, {}, draft)
        pdf = render_pdf(draft, tmp)
        n_pages = pdf_pages(pdf)
        print(f"attempt {attempt}: capacity={capacity} screens={len(selected)} pages={n_pages}", flush=True)
        if n_pages <= args.target_pages:
            if best is None or n_pages > best[0]:
                best = (n_pages, list(selected))
        if args.target_pages - 6 <= n_pages <= args.target_pages:
            break
        lines_per_page = total_lines / max(n_pages - 6, 1)
        capacity += int((args.target_pages - 2 - n_pages) * lines_per_page)

    if best is None:
        sys.exit("could not fit the page target")
    n_pages, selected = best
    draft = tmp / "draft.docx"
    groups, _ = build_document(selected, {}, draft)
    pdf = render_pdf(draft, tmp)
    heading_pages = locate_headings(pdf, groups)
    final = Path(args.out)
    groups, _ = build_document(selected, heading_pages, final)
    pdf = render_pdf(final, tmp)
    n_pages = pdf_pages(pdf)
    check = locate_headings(pdf, groups)
    if check != heading_pages:
        sys.exit("table of contents page numbers changed between passes")
    print(f"final pages={n_pages}")
    print("screens included:", len(selected))
    for f in selected:
        print("  ", f)


if __name__ == "__main__":
    main()
